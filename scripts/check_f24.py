"""Check specific cell F24 in Hindsight file."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl

file_path = "data/uploads/Hindsight-IBNR-to-Case-Ratio-Template-5-15-2025.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=False)

print(f"Sheet names: {wb.sheetnames}")
print()

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    # Check F24
    cell_f24 = sheet.cell(row=24, column=6)  # F = column 6
    print(f"Sheet '{sheet_name}' - F24:")
    print(f"  Value: {cell_f24.value}")
    print(f"  Type: {type(cell_f24.value)}")
    if cell_f24.value and str(cell_f24.value).startswith("="):
        print(f"  FORMULA FOUND: {cell_f24.value}")
    print()

wb.close()
