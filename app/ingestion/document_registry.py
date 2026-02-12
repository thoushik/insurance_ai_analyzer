"""
Insurance Document Intelligence Assistant
Ingestion Module - Document Registry

Maintains an inventory of all ingested documents with metadata.
"""

import json
from pathlib import Path
from datetime import datetime
from typing import Optional
from dataclasses import dataclass, field, asdict

from ..security import get_folder_guard, get_audit_logger
from .excel_parser import ExcelParser, WorkbookInfo
from .pdf_parser import PDFParser, PDFInfo


@dataclass
class DocumentEntry:
    """A single document entry in the registry."""
    id: str
    filename: str
    filepath: str
    doc_type: str  # "excel" or "pdf"
    ingested_at: str
    file_size: int
    metadata: dict = field(default_factory=dict)
    
    def to_dict(self) -> dict:
        return asdict(self)


class DocumentRegistry:
    """
    Maintains an inventory of ingested documents.
    
    Features:
    - Document tracking with metadata
    - Parsed data caching
    - Quick lookup for chat context
    """
    
    def __init__(self):
        self.guard = get_folder_guard()
        self.logger = get_audit_logger()
        
        self.cache_file = self.guard.cache_dir / "document_registry.json"
        self.documents: dict[str, DocumentEntry] = {}
        self.parsed_data: dict[str, WorkbookInfo | PDFInfo] = {}
        
        # Parsers
        self.excel_parser = ExcelParser()
        self.pdf_parser = PDFParser()
        
        # Load existing registry
        self._load_registry()
    
    def _load_registry(self):
        """Load registry from cache file."""
        if self.cache_file.exists():
            try:
                data = json.loads(self.cache_file.read_text(encoding="utf-8"))
                for doc_id, doc_data in data.items():
                    self.documents[doc_id] = DocumentEntry(**doc_data)
            except Exception as e:
                self.logger.log(
                    "registry_load_error",
                    "ingestion",
                    {"error": str(e)},
                    status="error"
                )
    
    def _save_registry(self):
        """Save registry to cache file."""
        data = {
            doc_id: doc.to_dict()
            for doc_id, doc in self.documents.items()
        }
        self.guard.safe_write(
            self.cache_file,
            json.dumps(data, indent=2, ensure_ascii=False)
        )
    
    def _generate_id(self, filename: str) -> str:
        """Generate a unique ID for a document."""
        import hashlib
        timestamp = datetime.now().isoformat()
        return hashlib.md5(f"{filename}:{timestamp}".encode()).hexdigest()[:12]
    
    def ingest_file(self, filepath: str | Path) -> DocumentEntry:
        """
        Ingest a single file into the registry.
        
        Args:
            filepath: Path to the file
            
        Returns:
            DocumentEntry for the ingested file
        """
        safe_path = self.guard.validate_path(filepath)
        
        if not safe_path.exists():
            raise FileNotFoundError(f"File not found: {filepath}")
        
        # Check for duplicate by filename - return existing entry if found
        existing_entry = self.get_by_filename(safe_path.name)
        if existing_entry:
            self.logger.log(
                "document_skipped_duplicate",
                "ingestion",
                {"filename": safe_path.name, "existing_id": existing_entry.id}
            )
            return existing_entry
        
        # Determine file type
        suffix = safe_path.suffix.lower()
        if suffix in [".xlsx", ".xls"]:
            doc_type = "excel"
            parsed = self.excel_parser.parse(safe_path)
            metadata = parsed.to_dict()
        elif suffix == ".pdf":
            doc_type = "pdf"
            parsed = self.pdf_parser.parse(safe_path)
            metadata = parsed.to_dict()
        else:
            raise ValueError(f"Unsupported file type: {suffix}")
        
        # Create entry
        doc_id = self._generate_id(safe_path.name)
        entry = DocumentEntry(
            id=doc_id,
            filename=safe_path.name,
            filepath=str(safe_path),
            doc_type=doc_type,
            ingested_at=datetime.now().isoformat(),
            file_size=safe_path.stat().st_size,
            metadata=metadata
        )
        
        # Store
        self.documents[doc_id] = entry
        self.parsed_data[doc_id] = parsed
        
        self._save_registry()
        
        self.logger.log(
            "document_registered",
            "ingestion",
            {
                "doc_id": doc_id,
                "filename": safe_path.name,
                "type": doc_type
            }
        )
        
        return entry
    
    def ingest_folder(self, folder_path: str | Path) -> list[DocumentEntry]:
        """
        Ingest all supported files from a folder and remove missing ones.
        
        Args:
            folder_path: Path to the folder
            
        Returns:
            List of DocumentEntry objects
        """
        safe_path = self.guard.validate_path(folder_path)
        
        if not safe_path.is_dir():
            raise ValueError(f"Not a directory: {folder_path}")
        
        # 1. Sync: Remove "ghost" documents not in this folder
        self._sync_folder(safe_path)
        
        entries = []
        supported_extensions = [".xlsx", ".xls", ".pdf"]
        
        for file_path in safe_path.rglob("*"):
            if file_path.suffix.lower() in supported_extensions:
                try:
                    entry = self.ingest_file(file_path)
                    entries.append(entry)
                except Exception as e:
                    self.logger.log(
                        "ingestion_error",
                        "ingestion",
                        {
                            "file": str(file_path),
                            "error": str(e)
                        },
                        status="error"
                    )
        
        return entries

    def _sync_folder(self, folder_path: Path):
        """
        Remove registry entries that are not present in the folder.
        """
        to_remove = []
        folder_str = str(folder_path)
        
        # Identify missing files
        for doc_id, doc in self.documents.items():
            # Check if file belongs to this folder (or subfolder)
            if doc.filepath.startswith(folder_str):
                if not Path(doc.filepath).exists():
                    to_remove.append(doc_id)
        
        # Remove them
        for doc_id in to_remove:
            self._remove_document(doc_id)
            
        if to_remove:
            self.logger.log(
                "registry_sync_removed",
                "ingestion",
                {"count": len(to_remove), "removed_ids": to_remove}
            )
            self._save_registry()

    def _remove_document(self, doc_id: str):
        """Remove a document from registry and memory."""
        if doc_id in self.documents:
            del self.documents[doc_id]
        if doc_id in self.parsed_data:
            del self.parsed_data[doc_id]
    
    def get_document(self, doc_id: str) -> Optional[DocumentEntry]:
        """Get a document by ID."""
        return self.documents.get(doc_id)
    
    def get_by_filename(self, filename: str) -> Optional[DocumentEntry]:
        """Get a document by filename (for deduplication)."""
        for doc in self.documents.values():
            if doc.filename == filename:
                return doc
        return None
    
    def get_parsed_data(self, doc_id: str) -> Optional[WorkbookInfo | PDFInfo]:
        """Get parsed data for a document."""
        if doc_id in self.parsed_data:
            return self.parsed_data[doc_id]
        
        # Re-parse if not in memory
        entry = self.get_document(doc_id)
        if entry:
            if entry.doc_type == "excel":
                parsed = self.excel_parser.parse(entry.filepath)
            else:
                parsed = self.pdf_parser.parse(entry.filepath)
            
            self.parsed_data[doc_id] = parsed
            return parsed
        
        return None
    
    def list_documents(self) -> list[dict]:
        """List all documents in the registry."""
        return [
            {
                "id": doc.id,
                "filename": doc.filename,
                "type": doc.doc_type,
                "ingested_at": doc.ingested_at,
                "file_size": doc.file_size
            }
            for doc in self.documents.values()
        ]
    
    def get_document_summary(self, doc_id: str) -> dict:
        """Get a summary of a document."""
        entry = self.get_document(doc_id)
        if not entry:
            return {"error": "Document not found"}
        
        return {
            "id": entry.id,
            "filename": entry.filename,
            "type": entry.doc_type,
            "ingested_at": entry.ingested_at,
            "metadata": entry.metadata
        }
    
    def get_all_context(self) -> str:
        """
        Get context string for all documents for LLM.
        """
        context_parts = []
        
        for doc_id, entry in self.documents.items():
            context_parts.append(f"\n## Document: {entry.filename}")
            context_parts.append(f"Type: {entry.doc_type}")
            
            if entry.doc_type == "excel":
                meta = entry.metadata
                context_parts.append(f"Sheets: {meta.get('sheet_count', 0)}")
                for sheet in meta.get("sheets", []):
                    context_parts.append(
                        f"  - {sheet['name']} ({sheet['purpose']}): "
                        f"{sheet['row_count']} rows, {sheet['cells_with_formulas']} formulas"
                    )
                    if sheet.get("sample_formulas"):
                        context_parts.append("    Sample formulas:")
                        for f in sheet["sample_formulas"][:3]:
                            context_parts.append(f"      {f['cell']}: {f['formula']}")
            
            elif entry.doc_type == "pdf":
                meta = entry.metadata
                context_parts.append(f"Pages: {meta.get('page_count', 0)}")
                context_parts.append(f"Words: {meta.get('total_words', 0)}")
                context_parts.append(f"Tables: {meta.get('total_tables', 0)}")
        
        return "\n".join(context_parts)
    
    def clear_registry(self):
        """Clear all documents from the registry."""
        self.documents.clear()
        self.parsed_data.clear()
        self._save_registry()
    
    def get_cell_info(self, doc_id: str, cell_ref: str, sheet_name: str = None) -> dict:
        """
        Get information about a specific cell.
        
        Args:
            doc_id: Document ID
            cell_ref: Cell reference like 'F24' or 'Sheet1!F24'
            sheet_name: Optional sheet name
            
        Returns:
            Dictionary with cell info including formula and value
        """
        entry = self.get_document(doc_id)
        if not entry or entry.doc_type != "excel":
            return {"error": "Document not found or not an Excel file"}
        
        import openpyxl
        from openpyxl.utils import column_index_from_string
        import re
        
        # Parse cell reference
        if "!" in cell_ref:
            sheet_name, cell_ref = cell_ref.split("!", 1)
        
        # Parse column and row from cell reference
        match = re.match(r"([A-Za-z]+)(\d+)", cell_ref)
        if not match:
            return {"error": f"Invalid cell reference: {cell_ref}"}
        
        col_letter = match.group(1).upper()
        row = int(match.group(2))
        col = column_index_from_string(col_letter)
        
        try:
            wb = openpyxl.load_workbook(entry.filepath, data_only=False)
            wb_values = openpyxl.load_workbook(entry.filepath, data_only=True)
            
            # Find the sheet
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    return {"error": f"Sheet '{sheet_name}' not found"}
                target_sheets = [sheet_name]
            else:
                target_sheets = wb.sheetnames
            
            results = []
            for sname in target_sheets:
                sheet = wb[sname]
                sheet_values = wb_values[sname]
                
                cell = sheet.cell(row=row, column=col)
                cell_val = sheet_values.cell(row=row, column=col)
                
                if cell.value is not None or cell_val.value is not None:
                    results.append({
                        "sheet": sname,
                        "cell": cell_ref.upper(),
                        "formula": cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None,
                        "value": cell_val.value,
                        "raw_value": cell.value
                    })
            
            wb.close()
            wb_values.close()
            
            if results:
                return {"found": True, "cells": results}
            else:
                return {"found": False, "message": f"Cell {cell_ref} is empty in all sheets"}
                
        except Exception as e:
            return {"error": str(e)}
    
    def get_all_formulas_for_document(self, doc_id: str) -> list:
        """
        Get all formulas from an Excel document.
        
        Args:
            doc_id: Document ID
            
        Returns:
            List of all formulas with their locations and values
        """
        entry = self.get_document(doc_id)
        if not entry or entry.doc_type != "excel":
            return []
        
        return self.excel_parser.extract_formulas(entry.filepath)


# Global instance
_registry = None

def get_document_registry() -> DocumentRegistry:
    """Get the global DocumentRegistry instance."""
    global _registry
    if _registry is None:
        _registry = DocumentRegistry()
    return _registry
